#!/usr/bin/env python3
"""Dump Xindus customer.xlsx to JSON for CRM ERP backfill (stdout)."""
import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

XLSX = Path(sys.argv[1] if len(sys.argv) > 1 else "/Users/apple/Downloads/Xindus customer.xlsx")


def json_val(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return str(v) if abs(v) > 2**53 else v
    if isinstance(v, float):
        if v.is_integer() and abs(v) > 2**53:
            return str(int(v))
        return v
    return str(v)


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = []
    for raw in it:
        row = {}
        empty = True
        for i, h in enumerate(headers):
            if not h or i >= len(raw):
                continue
            v = raw[i]
            if v is None or v == "":
                continue
            empty = False
            row[h] = json_val(v)
        if not empty:
            rows.append(row)
    wb.close()
    json.dump(rows, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
